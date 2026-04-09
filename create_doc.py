from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    if level == 1:
        heading_format = heading.runs[0]
        heading_format.font.color.rgb = RGBColor(46, 80, 144)
    return heading

def shade_cell(cell, color):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._element.get_or_add_tcPr().append(shading_elm)

doc = Document()

# Set default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# Title Page
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Project Architecture Overview')
run.font.size = Pt(36)
run.font.bold = True
run.font.color.rgb = RGBColor(46, 80, 144)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('NYC Air Quality Field Monitoring System')
run.font.size = Pt(16)
run.font.italic = True

doc.add_page_break()

# Table of Contents
add_heading(doc, "Table of Contents", level=1)
toc_items = [
    "Executive Summary",
    "System Overview",
    "System Architecture",
    "Core Scripts",
    "Data Flow",
    "Key Data Sources",
    "Technology Stack",
    "Key Interconnections",
    "Quick Reference",
    "For New Team Members"
]
for item in toc_items:
    p = doc.add_paragraph(item, style='List Bullet')

doc.add_page_break()

# Executive Summary
add_heading(doc, "Executive Summary", level=1)
doc.add_paragraph(
    "A real-time field monitoring platform for air quality research where collectors walk predetermined routes with GPS trackers. "
    "The system automates walk scheduling based on weather forecasts and collector availability, then provides real-time monitoring dashboards for the team."
)

# System Overview
add_heading(doc, "System Overview", level=1)
doc.add_paragraph("What: Real-time field monitoring for air quality research", style='List Bullet')
doc.add_paragraph("Who: NYC field collectors walking predetermined routes", style='List Bullet')
doc.add_paragraph("How: Automated scheduling using Claude AI + weather + availability data", style='List Bullet')
doc.add_paragraph("Where: Deployed on Fly.io with persistent data storage", style='List Bullet')

# System Architecture
add_heading(doc, "System Architecture", level=1)
arch_text = """FIELD LAYER
├─ GPS Trackers (Collector phones)
└─ Google Drive (Walk data uploads)
       ↓
    [serve.py] ← Central hub
    ↙         ↘
walk_scheduler.py    build_dashboard.py
    ↓                    ↓
Recommendations    Dashboard HTML"""

p = doc.add_paragraph(arch_text)
p_format = p.paragraph_format
p_format.left_indent = Inches(0.25)
for run in p.runs:
    run.font.name = 'Courier New'
    run.font.size = Pt(10)

# Core Scripts
add_heading(doc, "Core Scripts", level=1)

# serve.py
add_heading(doc, "1. serve.py — HTTP Server & Data Hub", level=2)
p = doc.add_paragraph()
p.add_run("Purpose: ").bold = True
p.add_run("Central coordinator receiving GPS updates, polling data sources, and serving dashboards")

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Key Responsibilities:").bold = True
doc.add_paragraph("Polls Google Drive every 60 seconds for new walk data files", style='List Bullet')
doc.add_paragraph("Maintains Walks_Log.txt as single source of truth", style='List Bullet')
doc.add_paragraph("Serves HTML dashboards to web browsers", style='List Bullet')
doc.add_paragraph("Triggers scheduler and dashboard rebuilds", style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Input: ").bold = True
p.add_run("GPS updates, Google Drive files, user requests")

p = doc.add_paragraph()
p.add_run("Output: ").bold = True
p.add_run("JSON data, HTML pages, walk logs")

p = doc.add_paragraph()
p.add_run("Port: ").bold = True
p.add_run("8765 (IDE), 8080 (production)")

# walk_scheduler.py
add_heading(doc, "2. walk_scheduler.py — AI-Powered Scheduling Engine", level=2)
p = doc.add_paragraph()
p.add_run("Purpose: ").bold = True
p.add_run("Recommends optimal walks using Claude AI based on weather and availability")

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Key Responsibilities:").bold = True
doc.add_paragraph("Reads weekly weather forecast PDFs", style='List Bullet')
doc.add_paragraph("Reads collector availability schedules (Excel + PDFs)", style='List Bullet')
doc.add_paragraph("Reads walk history (Walks_Log.txt)", style='List Bullet')
doc.add_paragraph("Uses Claude API to generate intelligent recommendations", style='List Bullet')
doc.add_paragraph("Outputs schedule_output.json with suggestions", style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Input: ").bold = True
p.add_run("Forecast PDFs, Excel schedules, walk history")

p = doc.add_paragraph()
p.add_run("Output: ").bold = True
p.add_run("schedule_output.json (AI recommendations)")

p = doc.add_paragraph()
p.add_run("Dependencies: ").bold = True
p.add_run("Claude API, pdfplumber, openpyxl")

# build_dashboard.py
add_heading(doc, "3. build_dashboard.py — Dashboard Generator", level=2)
p = doc.add_paragraph()
p.add_run("Purpose: ").bold = True
p.add_run("Creates interactive HTML dashboard for real-time monitoring")

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Key Responsibilities:").bold = True
doc.add_paragraph("Parses route KML files into JSON", style='List Bullet')
doc.add_paragraph("Combines GPS position data with route overlays", style='List Bullet')
doc.add_paragraph("Builds walk log viewer", style='List Bullet')
doc.add_paragraph("Generates Leaflet.js interactive map with controls", style='List Bullet')
doc.add_paragraph("Embeds live data snapshots", style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Input: ").bold = True
p.add_run("routes_data.json, GPS logs, walk history")

p = doc.add_paragraph()
p.add_run("Output: ").bold = True
p.add_run("dashboard.html (dark-themed, interactive map)")

# build_collector_map.py
add_heading(doc, "4. build_collector_map.py — Collector Statistics Map", level=2)
p = doc.add_paragraph()
p.add_run("Purpose: ").bold = True
p.add_run("Displays collector home locations and activity statistics")

p = doc.add_paragraph()
p.add_run("Output: ").bold = True
p.add_run("collector_map.html (split-screen with pins and stats)")

# transit_matrix.py
add_heading(doc, "5. transit_matrix.py — Transit Analysis", level=2)
p = doc.add_paragraph()
p.add_run("Purpose: ").bold = True
p.add_run("Analyzes subway connectivity between routes")

p = doc.add_paragraph()
p.add_run("Output: ").bold = True
p.add_run("transit_matrix.json (subway connectivity data)")

doc.add_page_break()

# Data Flow
add_heading(doc, "Data Flow", level=1)

dataflow_text = """INPUT SOURCES:
├─ GPS Trackers → serve.py (real-time updates)
├─ Google Drive → serve.py (walk completion files)
├─ Forecast PDFs → walk_scheduler.py (weather data)
├─ Collector Schedules (Excel) → walk_scheduler.py (availability)
├─ Route KML Files → build_dashboard.py (map geometry)
└─ NYC Transit GTFS → transit_matrix.py (subway reference)

PROCESSING:
├─ serve.py: Aggregates updates, serves dashboards
└─ walk_scheduler.py: Generates AI-powered recommendations

OUTPUT/STORAGE:
├─ dashboard.html → Real-time team monitoring
├─ schedule_output.json → Schedule recommendations
├─ collector_map.html → Location analytics
├─ routes_data.json → Embedded route data
└─ Walks_Log.txt → Master walk completion log"""

p = doc.add_paragraph(dataflow_text)
p_format = p.paragraph_format
p_format.left_indent = Inches(0.25)
for run in p.runs:
    run.font.name = 'Courier New'
    run.font.size = Pt(9)

# Key Data Sources Table
add_heading(doc, "Key Data Sources", level=1)

table = doc.add_table(rows=7, cols=2)
table.style = 'Light Grid Accent 1'

header_cells = table.rows[0].cells
header_cells[0].text = 'File'
header_cells[1].text = 'Purpose'

shade_cell(header_cells[0], "D5E8F0")
shade_cell(header_cells[1], "D5E8F0")

data_sources = [
    ("Walks_Log.txt", "Master log of completed walks"),
    ("Preferred_Routes.xlsx", "Route preferences per collector"),
    ("Forecast/", "Weekly weather forecasts (PDFs)"),
    ("Route_KMLs/", "Geographic route data (4 NYC boroughs)"),
    ("Collector_Schedule/", "Individual availability schedules"),
    ("Subway_gtfs/", "NYC MTA transit reference data"),
]

for i, (file, purpose) in enumerate(data_sources, 1):
    row_cells = table.rows[i].cells
    row_cells[0].text = file
    row_cells[1].text = purpose

# Technology Stack Table
add_heading(doc, "Technology Stack", level=1)

table = doc.add_table(rows=8, cols=2)
table.style = 'Light Grid Accent 1'

header_cells = table.rows[0].cells
header_cells[0].text = 'Layer'
header_cells[1].text = 'Technology'

shade_cell(header_cells[0], "D5E8F0")
shade_cell(header_cells[1], "D5E8F0")

tech_stack = [
    ("Backend", "Python 3 (HTTP server, data processing)"),
    ("Frontend", "HTML/CSS/JavaScript (Leaflet.js maps)"),
    ("AI", "Claude API (scheduling optimization)"),
    ("Data Sources", "Google Drive API, Excel, KML, PDF"),
    ("Hosting", "Fly.io (persistent /data volume)"),
    ("Dependencies", "anthropic, google-api-client, pdfplumber, openpyxl, pandas"),
    ("Extra", ""),
]

for i, (layer, tech) in enumerate(tech_stack[:7], 1):
    row_cells = table.rows[i].cells
    row_cells[0].text = layer
    row_cells[1].text = tech

doc.add_page_break()

# Key Interconnections
add_heading(doc, "Key Interconnections", level=1)

table = doc.add_table(rows=9, cols=3)
table.style = 'Light Grid Accent 1'

header_cells = table.rows[0].cells
header_cells[0].text = 'Component A'
header_cells[1].text = 'Component B'
header_cells[2].text = 'Connection'

for cell in header_cells:
    shade_cell(cell, "D5E8F0")

interconnections = [
    ("serve.py", "walk_scheduler.py", "Shares walk logs and configuration"),
    ("serve.py", "build_dashboard.py", "Triggered on GPS/data updates"),
    ("walk_scheduler.py", "Forecast PDFs", "Reads weather for optimization"),
    ("walk_scheduler.py", "Excel schedules", "Reads collector availability"),
    ("All outputs", "Fly.io /data", "Persistent storage and sync"),
    ("GPS trackers", "serve.py", "Real-time position updates"),
    ("Google Drive", "serve.py", "Polling (60-second intervals)"),
    ("", "", ""),
]

for i, (compA, compB, conn) in enumerate(interconnections, 1):
    row_cells = table.rows[i].cells
    row_cells[0].text = compA
    row_cells[1].text = compB
    row_cells[2].text = conn

# Quick Reference
add_heading(doc, "Quick Reference: Script Execution Flow", level=1)

flow_steps = [
    ("Startup", "serve.py initializes, listens on port 8080"),
    ("GPS Updates", "Field trackers send positions to serve.py"),
    ("Data Polling", "serve.py checks Google Drive every 60 seconds"),
    ("Scheduling", "walk_scheduler.py runs, reads data, uses Claude API"),
    ("Dashboard Generation", "build_dashboard.py regenerates HTML when data updates"),
    ("Monitoring", "Teams access dashboard.html in browser"),
    ("Storage", "All outputs saved to Fly.io persistent volume"),
]

for step, desc in flow_steps:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(step + ": ").bold = True
    p.add_run(desc)

# For New Team Members
add_heading(doc, "For New Team Members", level=1)

doc.add_paragraph("To understand this system:")
doc.add_paragraph("Start with: System Overview section", style='List Bullet')
doc.add_paragraph("Then read: Core Scripts section", style='List Bullet')
doc.add_paragraph("Understand: Data Flow diagram", style='List Bullet')
doc.add_paragraph("Reference: Technology Stack & Interconnections", style='List Bullet')

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("The entire system revolves around ")
run = p.add_run("serve.py")
run.bold = True
p.add_run(" as the central hub, with specialized workers triggered by data updates.")

# Save document
doc.save('Architecture_Overview.docx')
print("Document created successfully: Architecture_Overview.docx")
